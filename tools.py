from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from config import tavily_client, OUTPUT_DIR
import datetime
import pandas as pd
import io
import os
import rag
import openpyxl



#联网工具
def web_search(query: str) -> str:
    print(f"[系统日志] Agent 正在联网搜索: {query}")
    try:
        # 调用 Tavily 搜索，限制返回 3 条结果以节省 token
        response = tavily_client.search(query=query, search_depth="basic", max_results=3)
        
        # 把搜索结果拼成一段文本
        results_text = ""
        for idx, result in enumerate(response.get("results", []), 1):
            results_text += f"结果 {idx}:\n标题: {result['title']}\n内容: {result['content']}\n\n"
            
        return results_text if results_text else "未找到相关搜索结果。"
    except Exception as e:
        return f"搜索失败，错误信息: {str(e)}"

#word文档工具
def generate_word(topic: str, content: str) -> str:
    print(f"[系统日志] Agent 正在生成 Word 文档: {topic}")
    try:
        doc = Document()
        doc.add_heading(topic, level=0)
        
        # 简单解析一下换行，写入段落
        paragraphs = content.split("\n")
        for p in paragraphs:
            if p.strip():
                doc.add_paragraph(p.strip())
                
        # 生成安全的文件名（防止非法字符）
        safe_topic = "".join([c for c in topic if c.isalnum() or c in " _-"]).strip().replace(" ", "_")
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{safe_topic}_{timestamp}.docx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        
        doc.save(filepath)
        return f"文档生成成功，已保存为：{filename}"
    except Exception as e:
        return f"文档生成失败，错误信息: {str(e)}"

#RAG检索工具
def local_knowledge_search(query: str) -> str:
    print(f"[系统日志] Agent 正在检索本地知识库: {query}")
    try:
        # 调用混合检索函数（向量+BM25+Rerank）
        results = rag.search_similar(query, top_k=3)
        if not results:
            return "本地知识库中没有找到相关内容。"
        
        text = ""
        for i, res in enumerate(results, 1):
            text += f"【本地文档片段 {i}】来源: {res['source']}\n{res['text']}\n\n"
        return text
    except Exception as e:
        return f"本地检索失败: {str(e)}"

#表格生成工具s
def generate_excel(topic: str, table_data: str) -> str:
    print(f"[系统日志] Agent 正在生成 Excel 表格: {topic}")
    try:
        # 1. 容错清洗：防止大模型自作主张加上 Markdown 的 ```csv 标记
        table_data = table_data.strip()
        if table_data.startswith("```csv"):
            table_data = table_data[6:]
        if table_data.startswith("```"):
            table_data = table_data[3:]
        if table_data.endswith("```"):
            table_data = table_data[:-3]
        table_data = table_data.strip()
        
        # 2. 使用 pandas 解析 CSV 格式字符串
        df = pd.read_csv(io.StringIO(table_data))
        
        # 3. 生成安全的文件名
        safe_topic = "".join([c for c in topic if c.isalnum() or c in " _-"]).strip().replace(" ", "_")
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{safe_topic}_{timestamp}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        
        # 先把数据导出为最基础的 Excel
        df.to_excel(filepath, index=False)
        
        # 4. 美化 Excel (使用 openpyxl 重新打开并操作)
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        
        # 定义样式
        header_font = Font(name='微软雅黑', bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="409EFF", end_color="008500", fill_type="solid") # 蓝色表头
        center_alignment = Alignment(horizontal="center", vertical="center")
        left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style='thin', color="D3D3D3"), 
                            right=Side(style='thin', color="D3D3D3"), 
                            top=Side(style='thin', color="D3D3D3"), 
                            bottom=Side(style='thin', color="D3D3D3"))
        
        # 遍历所有列进行美化
        for col_idx, col_name in enumerate(df.columns, 1):
            # 美化表头 (第一行)
            header_cell = ws.cell(row=1, column=col_idx)
            header_cell.font = header_font
            header_cell.fill = header_fill
            header_cell.alignment = center_alignment
            header_cell.border = thin_border
            
            # 美化数据行 (第二行到最后一行)
            max_length = len(str(col_name))
            for row_idx in range(2, len(df) + 2):
                data_cell = ws.cell(row=row_idx, column=col_idx)
                data_cell.font = Font(name='微软雅黑', size=10)
                data_cell.alignment = left_alignment
                data_cell.border = thin_border
                
                # 计算列宽所需的最大字符数
                cell_value = str(data_cell.value) if data_cell.value is not None else ""
                max_length = max(max_length, len(cell_value))
                
            # 设置列宽 (中文字符较宽，乘以 1.5 的系数，并限制最大宽度 50)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max_length * 1.5 + 5, 50)
            
        wb.save(filepath)
        
        return f"表格生成成功，已保存为：{filepath}"
    except Exception as e:
        print(f"Excel 生成错误详情: {str(e)}")
        return f"表格生成失败，错误信息: {str(e)}"

available_functions = {
    "web_search": web_search,
    "generate_word": generate_word,
    "local_knowledge_search": local_knowledge_search,
    "generate_excel": generate_excel
}